"""Собирает упоминания.xlsx из данные/<год>/*.csv.

Вкладки: «Все упоминания», затем по одной на каждый год (1890, 1891, …)
и «Выпуски <год>» — журнал просмотренных номеров.
Запуск из корня репозитория:  py -3.13 инструменты/собрать_excel.py
"""
import csv
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "данные"
OUT = ROOT / "упоминания.xlsx"

WIDTHS = {"id": 10, "год": 6, "раздел": 7, "объект": 16, "тема": 40, "выпуск": 7, "дата": 11,
          "файл": 18, "стр_pdf": 8, "стр_газеты": 10, "рубрика": 30, "публикация": 10,
          "лица": 35, "цитата": 70, "смысл": 50, "фрагменты": 40, "результат": 60}
HEAD_FILL = PatternFill("solid", fgColor="DDE4EE")
OBJ_FILL = {"Ветка": "FFF4D6", "Ветковская волость": "FFF4D6", "Имение Ветка": "E8F1E0",
            "Грошиков": "F6E1E1", "Рубанов": "E3E7F6"}


def read(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    if not rows:
        return
    cols = list(rows[0].keys())
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for r in rows:
        ws.append([int(v) if v.isdigit() and k in ("год", "раздел", "выпуск") else v for k, v in r.items()])
        fill = OBJ_FILL.get(r.get("объект", ""))
        if fill:
            ws.cell(ws.max_row, cols.index("объект") + 1).fill = PatternFill("solid", fgColor=fill)
        # первый фрагмент — кликабельная ссылка (относительный путь, работает рядом с папкой фрагменты/)
        if r.get("фрагменты"):
            c = ws.cell(ws.max_row, cols.index("фрагменты") + 1)
            c.hyperlink = r["фрагменты"].split(";")[0].strip()
            c.font = Font(color="1F4E99", underline="single")
    for i, k in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(k, 14)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def main():
    years = sorted(p.name for p in DATA.iterdir() if p.is_dir() and p.name.isdigit())
    wb = Workbook()
    wb.remove(wb.active)
    mentions = {y: read(DATA / y / "упоминания.csv") for y in years if (DATA / y / "упоминания.csv").exists()}
    sheet(wb, "Все упоминания", [r for y in years for r in mentions.get(y, [])])
    for y in years:
        sheet(wb, y, mentions.get(y, []))
    for y in years:
        if (DATA / y / "выпуски.csv").exists():
            sheet(wb, f"Выпуски {y}", read(DATA / y / "выпуски.csv"))
    wb.save(OUT)
    print(f"{OUT.name}: годы {', '.join(years)}; упоминаний {sum(map(len, mentions.values()))}")


if __name__ == "__main__":
    main()
