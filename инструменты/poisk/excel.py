"""Excel в том же виде, что упоминания.xlsx (инструменты/собрать_excel.py в папке «могилевские ведомости»)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Колонки упоминания.csv — см. данные/СХЕМА.md
MENTION_FIELDS = ["id", "год", "раздел", "объект", "тема", "выпуск", "дата", "файл", "стр_pdf", "стр_газеты",
                  "рубрика", "публикация", "лица", "цитата", "смысл", "фрагменты"]
PAGES_FIELDS = ["файл", "стр_pdf", "выпуск", "дата", "стр_газеты", "результат"]

WIDTHS = {"id": 10, "год": 6, "раздел": 7, "объект": 16, "тема": 40, "выпуск": 7, "дата": 11,
          "файл": 18, "стр_pdf": 8, "стр_газеты": 10, "рубрика": 30, "публикация": 10,
          "лица": 35, "цитата": 70, "смысл": 50, "фрагменты": 40, "результат": 60, "статус": 14}
HEAD_FILL = PatternFill("solid", fgColor="DDE4EE")
OBJ_FILL = {"Ветка": "FFF4D6", "Ветковская волость": "FFF4D6", "Имение Ветка": "E8F1E0",
            "Грошиков": "F6E1E1", "Рубанов": "E3E7F6"}
STATUS_FILL = {"не проверено": "EEEEEE", "отклонено": "F2DCDB", "ошибка проверки": "F2DCDB"}
NUMERIC = ("год", "раздел", "выпуск")


def sheet(wb, title, rows, cols):
    ws = wb.create_sheet(title[:31])
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for r in rows:
        ws.append([int(v) if k in NUMERIC and str(v).isdigit() else v for k, v in ((k, r.get(k, "")) for k in cols)])
        fill = OBJ_FILL.get(r.get("объект", ""))
        if fill and "объект" in cols:
            ws.cell(ws.max_row, cols.index("объект") + 1).fill = PatternFill("solid", fgColor=fill)
        fill = STATUS_FILL.get(r.get("статус", ""))
        if fill and "статус" in cols:
            ws.cell(ws.max_row, cols.index("статус") + 1).fill = PatternFill("solid", fgColor=fill)
        # первый фрагмент — кликабельная ссылка (относительный путь, работает рядом с папкой фрагменты/)
        if r.get("фрагменты") and "фрагменты" in cols:
            c = ws.cell(ws.max_row, cols.index("фрагменты") + 1)
            c.hyperlink = r["фрагменты"].split(";")[0].strip()
            c.font = Font(color="1F4E99", underline="single")
    for i, k in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(k, 14)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def write(path, mentions, pages, with_status=False):
    """mentions — строки по MENTION_FIELDS (+ «статус»), pages — строки по PAGES_FIELDS с ключом «год»."""
    cols = MENTION_FIELDS + (["статус"] if with_status else [])
    years = sorted({str(r["год"]) for r in mentions} | {str(r["год"]) for r in pages})
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Все упоминания", mentions, cols)
    for y in years:
        sheet(wb, y, [r for r in mentions if str(r["год"]) == y], cols)
    for y in years:
        sheet(wb, f"Просмотрено {y}", [r for r in pages if str(r["год"]) == y], PAGES_FIELDS)
    wb.save(path)
